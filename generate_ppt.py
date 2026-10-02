import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    
    # 16:9 와이드 슬라이드 크기 설정 (13.333 x 7.5 인치)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # 테마 색상 정의 (애플 감성 톤앤매너)
    COLOR_BG = RGBColor(245, 245, 247)       # #F5F5F7 라이트 그레이
    COLOR_WHITE = RGBColor(255, 255, 255)   # #FFFFFF
    COLOR_DARK = RGBColor(29, 29, 31)       # #1D1D1F 메인 텍스트
    COLOR_SUB = RGBColor(134, 134, 139)     # #86868B 서브 텍스트
    COLOR_BLUE = RGBColor(0, 122, 255)      # #007AFF 애플 블루
    COLOR_GREEN = RGBColor(52, 199, 89)     # #34C759 에메랄드 그린
    COLOR_PURPLE = RGBColor(88, 86, 214)    # #5856D6 인디고 퍼플
    COLOR_AMBER = RGBColor(245, 158, 11)    # #F59E0B 앰버 오렌지
    COLOR_BORDER = RGBColor(229, 229, 234)  # #E5E5EA 카디두께 테두리
    
    FONT_NAME = "맑은 고딕"
    
    blank_layout = prs.slide_layouts[6]
    
    # 공통 배경 생성 헬퍼
    def add_bg(slide, color=COLOR_BG):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    # 공통 카드 쉐입 생성 헬퍼
    def add_card(slide, left, top, width, height, bg_color=COLOR_WHITE, border_color=COLOR_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1)
        else:
            card.line.fill.background()
        return card

    # 공통 헤더 헬퍼
    def add_header(slide, title_text, category_text="사용자 가이드 매뉴얼"):
        # 카테고리 태그
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10), Inches(0.4))
        tf_tag = tag_box.text_frame
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = category_text.upper()
        p_tag.font.name = FONT_NAME
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_BLUE
        
        # 메인 타이틀
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = FONT_NAME
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_DARK

    # =========================================================================
    # Slide 1: 표지 (Title Slide)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    add_bg(slide1, COLOR_BG)
    
    # 메인 표지 카드 컨테이너
    add_card(slide1, Inches(1.0), Inches(1.0), Inches(11.333), Inches(5.5), COLOR_WHITE, COLOR_BORDER)
    
    # 서비스 로고 아이콘 상징 박스
    logo_box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.6), Inches(1.8), Inches(1.0), Inches(1.0))
    logo_box.fill.solid()
    logo_box.fill.fore_color.rgb = COLOR_BLUE
    logo_box.line.fill.background()
    tf_logo = logo_box.text_frame
    p_logo = tf_logo.paragraphs[0]
    p_logo.text = "🎓"
    p_logo.alignment = PP_ALIGN.CENTER
    p_logo.font.size = Pt(32)
    
    # 타이틀
    t_box = slide1.shapes.add_textbox(Inches(1.6), Inches(3.0), Inches(10.0), Inches(2.0))
    tf = t_box.text_frame
    tf.word_wrap = True
    
    p1 = tf.paragraphs[0]
    p1.text = "인근 고등학교 수시 합격 DB"
    p1.font.name = FONT_NAME
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_DARK
    
    p2 = tf.add_paragraph()
    p2.text = "선생님을 위한 맞춤형 입시 상담 및 통합 데이터 관리 시스템 사용 가이드"
    p2.font.name = FONT_NAME
    p2.font.size = Pt(18)
    p2.font.color.rgb = COLOR_SUB
    p2.space_before = Pt(12)

    # 하단 바
    footer_card = add_card(slide1, Inches(1.6), Inches(5.2), Inches(10.133), Inches(0.8), COLOR_BG, COLOR_BORDER)
    tf_foot = footer_card.text_frame
    p_foot = tf_foot.paragraphs[0]
    p_foot.text = "  📌 인증 기반 14명 상담진 공동 데이터베이스  |  구글 시트 DB 실시간 자동 연동"
    p_foot.font.name = FONT_NAME
    p_foot.font.size = Pt(13)
    p_foot.font.bold = True
    p_foot.font.color.rgb = COLOR_DARK
    p_foot.alignment = PP_ALIGN.LEFT

    # =========================================================================
    # Slide 2: 주요 기능 및 서비스 개요 (Overview)
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_bg(slide2)
    add_header(slide2, "1. 시스템 핵심 개요 및 주요 기능 3가지")

    card_width = Inches(3.64)
    card_height = Inches(5.2)
    
    # 기능 1: 실시간 대시보드
    add_card(slide2, Inches(0.8), Inches(1.6), card_width, card_height)
    tb = slide2.shapes.add_textbox(Inches(1.0), Inches(1.8), card_width - Inches(0.4), card_height - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "📊 1. 보기 탭 (대시보드)"
    p.font.name = FONT_NAME
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE
    
    items1 = [
        "• 관리 대상 고교 및 수시 합격 실적 한눈에 파악",
        "• 학년도(숫자 정규화) 및 학교명 실시간 검색 필터링",
        "• 학교 카드 렌더링: [전체 학생수] & [고교 유형] 배지 바인딩",
        "• 학교 카드 탭 시 iOS 시트 스타일 세부 모달 오픈"
    ]
    for it in items1:
        p = tf.add_paragraph()
        p.text = it
        p.font.name = FONT_NAME
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_DARK
        p.space_before = Pt(14)

    # 기능 2: 중복 합격 다중 입력
    add_card(slide2, Inches(4.84), Inches(1.6), card_width, card_height)
    tb = slide2.shapes.add_textbox(Inches(5.04), Inches(1.8), card_width - Inches(0.4), card_height - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "📝 2. 수시 합격 다중 입력"
    p.font.name = FONT_NAME
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_GREEN
    
    items2 = [
        "• 한 학생의 여러 대학/학과 중복 합격 실적 동시 등록",
        "• [+ 합격 항목 추가] 버튼으로 동적 입력 행 생성",
        "• 전교 등수, 내신 등급, 합격 대학/전형 1:1 매핑",
        "• 작성자 자동 반영으로 신뢰성 높은 정보 축적"
    ]
    for it in items2:
        p = tf.add_paragraph()
        p.text = it
        p.font.name = FONT_NAME
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_DARK
        p.space_before = Pt(14)

    # 기능 3: 학교 기본정보 & 프로그램
    add_card(slide2, Inches(8.88), Inches(1.6), card_width, card_height)
    tb = slide2.shapes.add_textbox(Inches(9.08), Inches(1.8), card_width - Inches(0.4), card_height - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "🏫 3. 학교정보 & 프로그램"
    p.font.name = FONT_NAME
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PURPLE
    
    items3 = [
        "• 학교 기본 정보 (고3 학생수, 고교 유형, 교과편성표)",
        "• 고3 학생수 텍스트 서식('350) 자동 보장",
        "• 특별 프로그램 및 우수 동아리 활동 카드 누적",
        "• 구글 드라이브 교과 편성표 외부 링크 연동"
    ]
    for it in items3:
        p = tf.add_paragraph()
        p.text = it
        p.font.name = FONT_NAME
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_DARK
        p.space_before = Pt(14)

    # =========================================================================
    # Slide 3: 1단계 - 교직원 인증 로그인 (Authentication)
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_bg(slide3)
    add_header(slide3, "2. [인증 및 접속] 교직원 로그인 방법")

    # 왼쪽 모달 UI 안내 카드
    add_card(slide3, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb = slide3.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.0), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "🔑 로그인 필수 입력 3가지 항목"
    p.font.name = FONT_NAME
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE
    
    login_steps = [
        ("1. 성명 (이름)", "상담진 목록에 등록된 선생님 성명 입력 (예: 분당러셀)"),
        ("2. 휴대폰 뒷자리 4개", "숫자 4자리 본인 확인용 (예: 1234)"),
        ("3. 사번 (또는 오늘 날짜)", "사번 입력 또는 Placeholder에 표시된 오늘 날짜(YYYYMMDD) 입력")
    ]
    for title, desc in login_steps:
        p = tf.add_paragraph()
        p.text = title
        p.font.name = FONT_NAME
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = COLOR_DARK
        p.space_before = Pt(14)
        
        p_sub = tf.add_paragraph()
        p_sub.text = desc
        p_sub.font.name = FONT_NAME
        p_sub.font.size = Pt(12)
        p_sub.font.color.rgb = COLOR_SUB

    # 오른쪽 프로세스 및 특이사항 카드
    add_card(slide3, Inches(6.8), Inches(1.6), Inches(5.733), Inches(5.2))
    tb = slide3.shapes.add_textbox(Inches(7.1), Inches(1.9), Inches(5.133), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "💡 로그인 접속 후 동작 프로세스"
    p.font.name = FONT_NAME
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_DARK
    
    features = [
        "✨ 로그인 성공 시 [입력 탭]으로 즉시 자동 전환!",
        "✨ 권한 분이: 14명 상담진 전원 인증 및 D열 관리자 특수 권한 부여",
        "✨ 세션 유지: 브라우저 세션 스토리지에 안전 보존 (자동 재로그인)",
        "✨ 작성자 자동 매핑: 데이터 입력 시 로그인한 선생님 이름 및 사번 자동 기록"
    ]
    for feat in features:
        p = tf.add_paragraph()
        p.text = feat
        p.font.name = FONT_NAME
        p.font.size = Pt(14)
        p.font.color.rgb = COLOR_DARK
        p.space_before = Pt(16)

    # =========================================================================
    # Slide 4: 2단계 - 보기 탭 (대시보드 활용법)
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_bg(slide4)
    add_header(slide4, "3. [보기 탭] 대시보드 및 학교 카운트 현황")

    # 대시보드 요소 4분할
    box_w = Inches(5.7)
    box_h = Inches(2.4)
    
    # 1. 요약 통계 카드
    add_card(slide4, Inches(0.8), Inches(1.6), box_w, box_h)
    tb = slide4.shapes.add_textbox(Inches(1.0), Inches(1.8), box_w - Inches(0.4), box_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "📈 상단 요약 통계 카드 4종"
    p.font.name = FONT_NAME
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE
    p2 = tf.add_paragraph()
    p2.text = "• 관리 대상 고교, 누적 수시 실적, 프로그램, 실시간 동기화 상태 표시\n• 구글 시트 변경 사항 자동 집계"
    p2.font.name = FONT_NAME
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_DARK
    p2.space_before = Pt(8)

    # 2. 필터링 및 검색
    add_card(slide4, Inches(6.8), Inches(1.6), box_w, box_h)
    tb = slide4.shapes.add_textbox(Inches(7.0), Inches(1.8), box_w - Inches(0.4), box_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🔍 학년도 필터 & 실시간 검색"
    p.font.name = FONT_NAME
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_GREEN
    p2 = tf.add_paragraph()
    p2.text = "• 학년도 선택('전체' 선택 시 모든 연도 100% 노출 보장)\n• 학교명 검색창으로 야탑고, 대진고 등 즉시 서치"
    p2.font.name = FONT_NAME
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_DARK
    p2.space_before = Pt(8)

    # 3. 학교별 종합 카드
    add_card(slide4, Inches(0.8), Inches(4.4), box_w, box_h)
    tb = slide4.shapes.add_textbox(Inches(1.0), Inches(4.6), box_w - Inches(0.4), box_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🎴 학교별 타일 카드 렌더링"
    p.font.name = FONT_NAME
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PURPLE
    p2 = tf.add_paragraph()
    p2.text = "• [전체 학생수] & [고교 유형] 배지 나란히 바인딩\n• 수시 합격 실적 건수 및 데이터 수집 진행도 프로그레스 바 제공"
    p2.font.name = FONT_NAME
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_DARK
    p2.space_before = Pt(8)

    # 4. 카드 탭 액션
    add_card(slide4, Inches(6.8), Inches(4.4), box_w, box_h)
    tb = slide4.shapes.add_textbox(Inches(7.0), Inches(4.6), box_w - Inches(0.4), box_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "👆 학교 카드 탭(클릭) 상세 조회"
    p.font.name = FONT_NAME
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_AMBER
    p2 = tf.add_paragraph()
    p2.text = "• 임의의 학교 타일을 누르면 세부 실적이 팝업 창으로 열림\n• 등수별 수시합격 내역 및 교과편성표 드라이브 링크 제공"
    p2.font.name = FONT_NAME
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_DARK
    p2.space_before = Pt(8)

    # =========================================================================
    # Slide 5: 3단계 - 학교 상세 정보 모달 (Detail Modal)
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_bg(slide5)
    add_header(slide5, "4. [학교 상세 모달] 세부 실적 및 편성표 열람")

    # 메인 모달 팝업 구조 설명
    add_card(slide5, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.2))
    tb = slide5.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(11.133), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "📱 iOS 시트 스타일 세부 모달 구성 요소"
    p.font.name = FONT_NAME
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE
    
    sections = [
        ("1️⃣ 학교 기본 정보 헤더", "학교명, 학년도 태그, 전체 학생 수, 고교 유형(일반고, 자사고 등) 1:1 시각적 바인딩"),
        ("2️⃣ 교과 편성표 드라이브 링크", "구글 드라이브 URL 등록 시 '교과 편성표 열람하기' 버튼이 동적 생성되어 바로 열람 가능"),
        ("3️⃣ 전교 등수별 수시 합격 실적", "전교 등수 순 정렬, 합격 대학, 합격 학과, 전형명, 내신 등급, 작성 상담진 명세 표시"),
        ("4️⃣ 특별 프로그램 및 동아리", "16~24px 여백과 애플 파스텔 배지가 적용된 특색 활동/우수 동아리 카드 리스트")
    ]
    for sec_title, sec_desc in sections:
        p = tf.add_paragraph()
        p.text = sec_title
        p.font.name = FONT_NAME
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = COLOR_DARK
        p.space_before = Pt(12)
        
        p_sub = tf.add_paragraph()
        p_sub.text = "    ➔ " + sec_desc
        p_sub.font.name = FONT_NAME
        p_sub.font.size = Pt(13)
        p_sub.font.color.rgb = COLOR_SUB

    # =========================================================================
    # Slide 6: 4단계 - 입력 탭 (3종 서브 탭 완벽 가이드)
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_bg(slide6)
    add_header(slide6, "5. [입력 탭] 데이터 축적 및 입력 서브 탭 3종")

    col_w = Inches(3.64)
    col_h = Inches(5.2)

    # 서브탭 1: 수시합격
    add_card(slide6, Inches(0.8), Inches(1.6), col_w, col_h)
    tb = slide6.shapes.add_textbox(Inches(1.0), Inches(1.8), col_w - Inches(0.4), col_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🥇 1. 수시 합격 입력"
    p.font.name = FONT_NAME
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_GREEN
    
    sub1 = [
        "• 연도, 학교명, 전교등수, 내신등급",
        "• [+ 합격 대학/학과 추가] 버튼으로 중복 합격 다중 입력",
        "• 한 번 제출 시 선택된 모든 대학 한꺼번에 누적 저장"
    ]
    for s in sub1:
        p = tf.add_paragraph()
        p.text = s
        p.font.name = FONT_NAME
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_DARK
        p.space_before = Pt(12)

    # 서브탭 2: 특별프로그램
    add_card(slide6, Inches(4.84), Inches(1.6), col_w, col_h)
    tb = slide6.shapes.add_textbox(Inches(5.04), Inches(1.8), col_w - Inches(0.4), col_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "✨ 2. 특별 프로그램"
    p.font.name = FONT_NAME
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_PURPLE
    
    sub2 = [
        "• 구분: 진행 특별 프로그램 vs 우수 동아리 활동",
        "• 연도, 학교명, 명칭, 주요 활동 내용 입력",
        "• 상담 시 시각적 특색 자료로 활용"
    ]
    for s in sub2:
        p = tf.add_paragraph()
        p.text = s
        p.font.name = FONT_NAME
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_DARK
        p.space_before = Pt(12)

    # 서브탭 3: 학교기본정보
    add_card(slide6, Inches(8.88), Inches(1.6), col_w, col_h)
    tb = slide6.shapes.add_textbox(Inches(9.08), Inches(1.8), col_w - Inches(0.4), col_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🏫 3. 학교 기본 정보"
    p.font.name = FONT_NAME
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE
    
    sub3 = [
        "• 학년도, 학교명, 고3 학생 수, 학교 유형, 교과편성표 링크",
        "• 구글 시트 A~G열 컬럼 1:1 매핑 (Upsert 갱신)",
        "• 고3 학생수 텍스트 서식('350) 자동 처리"
    ]
    for s in sub3:
        p = tf.add_paragraph()
        p.text = s
        p.font.name = FONT_NAME
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_DARK
        p.space_before = Pt(12)

    # =========================================================================
    # Slide 7: 5단계 - 구글 시트 연동 및 관리자 기능
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_bg(slide7)
    add_header(slide7, "6. [백엔드 연동] 구글 시트 DB & 관리자 설정")

    add_card(slide7, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.2))
    tb = slide7.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.1), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "📊 구글 시트 DB 구조 및 자동 생성"
    p.font.name = FONT_NAME
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE
    
    db_items = [
        "• '학교_기본정보' 시트 (A~G열 명세)",
        "  - A: 입력시간, B: 학교코드, C: 학년도, D: 학교명, E: 고3 학생수(텍스트), F: 학교 유형, G: 교과편성표 링크",
        "• '수시합격_입력' 시트 (A~I열 명세)",
        "• '특별프로그램_입력' 시트 (A~F열 명세)",
        "• 입력 시간(Timestamp) 서버 측 자동 세팅"
    ]
    for d in db_items:
        p = tf.add_paragraph()
        p.text = d
        p.font.name = FONT_NAME
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_DARK
        p.space_before = Pt(10)

    add_card(slide7, Inches(6.8), Inches(1.6), Inches(5.733), Inches(5.2))
    tb = slide7.shapes.add_textbox(Inches(7.1), Inches(1.9), Inches(5.133), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "⚙️ 관리자(D열) 연동 설정 및 동기화"
    p.font.name = FONT_NAME
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_AMBER
    
    adm_items = [
        "• 최고 관리자(admin) 계정 우측 상단 [연동 설정] 버튼 활성화",
        "• 구글 Apps Script 웹앱 URL 배포 후 연결",
        "• 데이터 수집 예외 처리(Try-Catch, Fallback) 적용",
        "• 실시간 수집 실패 시에도 로컬 캐시 데이터 100% 보장"
    ]
    for a in adm_items:
        p = tf.add_paragraph()
        p.text = a
        p.font.name = FONT_NAME
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_DARK
        p.space_before = Pt(14)

    # =========================================================================
    # Slide 8: Q&A 및 사용 유의사항 (FAQ & Tips)
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_bg(slide8)
    add_header(slide8, "7. [유의사항 및 FAQ] 데이터 관리 팁")

    add_card(slide8, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.2))
    tb = slide8.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(11.133), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "❓ 자주 묻는 질문 및 문제 해결"
    p.font.name = FONT_NAME
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE
    
    faqs = [
        ("Q. 대시보드에 카드가 표시되지 않을 때는 어떻게 하나요?",
         "A. 상단 '학년도 선택' 셀렉트박스가 '전체'로 설정되어 있는지 확인하시거나, 우측 상단 [새로고침] 버튼을 눌러 동기화해 주세요."),
        ("Q. 학교 기본 정보를 등록할 때 기존 데이터는 어떻게 되나요?",
         "A. 동일한 학교코드(예: 2026야탑고)나 [학년도+학교명] 조합이 있으면 기존 행이 자동으로 업데이트(Upsert)되고 없으면 신규 추가됩니다."),
        ("Q. 학생 수 데이터가 시트에서 지수로 바뀌거나 누락됩니다.",
         "A. 시스템에서 학생 수 저장 시 텍스트 서식('350)으로 자동 변환하여 시트에 넣으므로 누락 없이 안전하게 원본 유지됩니다.")
    ]
    for q, a in faqs:
        p = tf.add_paragraph()
        p.text = q
        p.font.name = FONT_NAME
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = COLOR_DARK
        p.space_before = Pt(14)
        
        p_ans = tf.add_paragraph()
        p_ans.text = a
        p_ans.font.name = FONT_NAME
        p_ans.font.size = Pt(12)
        p_ans.font.color.rgb = COLOR_SUB

    # 파일 저장
    output_path = r"C:\Users\A\.gemini\antigravity\scratch\highschool_consulting_app\인근고등학교_수시합격DB_사용법안내.pptx"
    prs.save(output_path)
    print(f"PPTX 파일이 성공적으로 생성되었습니다: {output_path}")

if __name__ == "__main__":
    create_presentation()
